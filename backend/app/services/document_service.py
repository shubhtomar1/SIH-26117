"""Upload → validate → store locally → extract text → SQLite. Treat content as UNTRUSTED DATA."""
import csv
import io
import json
import os
import uuid

from fastapi import UploadFile
from sqlalchemy.orm import Session

from app.core.config import ALLOWED_EXTENSIONS, settings
from app.core.logging import log
from app.models.document import Document

MAX_BYTES = settings.MAX_FILE_SIZE_MB * 1024 * 1024


def safe_filename(name: str) -> str:
    base = os.path.basename(name or "upload.bin")
    keep = "".join(c for c in base if c.isalnum() or c in ("-", "_", ".", " "))
    return (keep.strip() or "upload.bin")[:120]


def extract_text(path: str, ext: str) -> tuple[str, bool]:
    """Returns (text, needs_ocr). Never raises — returns '' on failure."""
    try:
        if ext in (".png", ".jpg", ".jpeg", ".webp", ".gif"):
            return "", True
        if ext in (".py", ".js", ".ts", ".c", ".cpp"):
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()[:20000], False
        if ext in (".txt", ".md", ".json", ".csv"):
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                raw = f.read()
            if ext == ".json":
                try:
                    return json.dumps(json.loads(raw))[:20000], False
                except ValueError:
                    pass
            if ext == ".csv":
                with open(path, newline="", encoding="utf-8", errors="ignore") as f:
                    rows = list(csv.reader(f))[:200]
                return "\n".join([", ".join(r) for r in rows])[:20000], False
            return raw[:20000], False
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(path)
            pages = reader.pages
            text = "\n".join([(p.extract_text() or "") for p in pages])[:20000]
            clean = text.strip()
            num_pages = max(len(pages), 1)

            has_images = any(len(p.images) > 0 for p in pages[:5])
            has_placeholder = any(ph in clean.lower() for ph in ("type your text", "sample scanned copy"))

            needs_ocr = False
            if len(clean) < 120:
                needs_ocr = True
            elif has_placeholder:
                needs_ocr = True
            elif has_images and (len(clean) / num_pages < 250):
                needs_ocr = True

            return text, needs_ocr
        if ext == ".docx":
            import docx

            doc = docx.Document(path)
            return "\n".join([p.text for p in doc.paragraphs])[:20000], False
        if ext == ".xlsx":
            import openpyxl

            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            out: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    out.append(", ".join([str(c) for c in row if c is not None]))
                    if len(out) > 500:
                        break
            return "\n".join(out)[:20000], False
        if ext == ".pptx":
            from pptx import Presentation

            prs = Presentation(path)
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        texts.append(shape.text)
            return "\n".join(texts)[:20000], False
    except Exception as exc:  # logged locally, never leaks
        log.warning(f"EXTRACT_FAILED path={path} err={exc}")
    return "", False


def is_scanned_pdf(path: str, content: str = "", needs_ocr: int | bool = 0) -> bool:
    """Determine whether a PDF is scanned or requires OCR extraction."""
    if needs_ocr:
        return True
    clean = (content or "").strip()
    if len(clean) < 120:
        return True
    lowered = clean.lower()
    if "type your text" in lowered or "sample scanned copy" in lowered:
        return True
    try:
        from pypdf import PdfReader
        r = PdfReader(path)
        pages = r.pages
        num_pages = max(len(pages), 1)
        has_imgs = any(len(p.images) > 0 for p in pages[:3])
        if has_imgs and (len(clean) / num_pages < 250):
            return True
    except Exception:
        pass
    return False


def ocr_pdf_pages(pdf_path: str, max_pages: int = 5) -> str:
    """Extract and transcribe scanned PDF pages via pypdfium2 / pypdf + Ollama vision model."""
    if not os.path.exists(pdf_path):
        return ""
    try:
        import base64
        import urllib.request
        from app.model_adapters.ollama_model import OLLAMA_URL, OLLAMA_VISION_MODEL

        pages_b64 = []

        # 1. High-fidelity rendering with pypdfium2
        try:
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(pdf_path)
            total = min(len(pdf), max_pages)
            for i in range(total):
                page = pdf[i]
                pil_img = page.render(scale=2).to_pil()
                buf = io.BytesIO()
                pil_img.save(buf, format="JPEG", quality=85)
                pages_b64.append(base64.b64encode(buf.getvalue()).decode())
        except Exception as exc:
            log.warning(f"PDFIUM_RENDER_FAILED path={pdf_path} err={exc}; falling back to pypdf images")

        # 2. Fallback to pypdf embedded image extraction
        if not pages_b64:
            try:
                from pypdf import PdfReader
                reader = PdfReader(pdf_path)
                for page in reader.pages[:max_pages]:
                    if page.images:
                        img_data = page.images[0].data
                        pages_b64.append(base64.b64encode(img_data).decode())
            except Exception as exc:
                log.warning(f"PYPDF_IMAGE_EXTRACT_FAILED path={pdf_path} err={exc}")

        if not pages_b64:
            log.warning(f"NO_PAGES_RENDERED_OR_EXTRACTED path={pdf_path}")
            return ""

        extracted = []
        for i, img_b64 in enumerate(pages_b64):
            payload = json.dumps({
                "model": OLLAMA_VISION_MODEL,
                "prompt": (
                    "Transcribe all readable text from this document page accurately, completely, and verbatim. "
                    "Preserve all names, designations, institutions, dates, reference numbers, addresses, headings, "
                    "certifications, paragraphs, signatures, and stamps. Do not omit any text."
                ),
                "images": [img_b64],
                "stream": False,
            }).encode()
            req = urllib.request.Request(
                f"{OLLAMA_URL}/api/generate",
                data=payload,
                headers={"Content-Type": "application/json"}
            )
            with urllib.request.urlopen(req, timeout=120) as resp:
                data = json.loads(resp.read().decode())
                text = (data.get("response") or "").strip()
                if text:
                    # Clean out boilerplate template artifacts if present
                    filtered = [
                        line for line in text.splitlines()
                        if line.strip().lower() not in (
                            "type your text",
                            "sample scanned copy:",
                            "sample scanned copy",
                        )
                    ]
                    cleaned = "\n".join(filtered).strip()
                    if cleaned:
                        extracted.append(f"--- Page {i + 1} ---\n{cleaned}")

        if extracted:
            return "\n\n".join(extracted)
    except Exception as exc:
        log.warning(f"SCANNED_PDF_OCR_FAILED path={pdf_path} err={exc}")
    return ""


def store_upload(db: Session, upload: UploadFile, department: str, category: str = "general") -> Document:
    ext = os.path.splitext(upload.filename or "")[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"UNSUPPORTED_FILE: extension '{ext}' not allowed.")
    data = upload.file.read()
    if len(data) > MAX_BYTES:
        raise ValueError(f"FILE_TOO_LARGE: max {settings.MAX_FILE_SIZE_MB} MB.")
    os.makedirs(settings.UPLOAD_DIR, exist_ok=True)
    doc_id = f"DOC-{uuid.uuid4().hex[:8].upper()}"
    stored = f"{doc_id}{ext}"
    path = os.path.abspath(os.path.join(settings.UPLOAD_DIR, stored))
    if not path.startswith(os.path.abspath(settings.UPLOAD_DIR)):  # path traversal guard
        raise ValueError("INVALID_FILE: unsafe path.")
    with open(path, "wb") as f:
        f.write(data)
    text, needs_ocr = extract_text(path, ext)
    if ext == ".pdf" and needs_ocr:
        try:
            scanned_text = ocr_pdf_pages(path, max_pages=3)
            if scanned_text:
                text = scanned_text
                needs_ocr = False
        except Exception as exc:
            log.warning(f"STORE_UPLOAD_OCR_FAILED path={path} err={exc}")
    doc = Document(
        id=doc_id,
        filename=safe_filename(upload.filename or stored),
        path=path,
        type=ext.lstrip("."),
        department=department,
        category=category,
        content=text,
        needs_ocr=1 if needs_ocr else 0,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    log.info(f"FILE_UPLOADED doc={doc_id} user_dept={department} ocr={needs_ocr}")
    return doc


def ingest_text(db: Session, filename: str, content: str, department: str, category: str = "general") -> Document:
    os.makedirs(settings.KNOWLEDGE_DIR, exist_ok=True)
    doc_id = f"DOC-{uuid.uuid4().hex[:8].upper()}"
    path = os.path.abspath(os.path.join(settings.KNOWLEDGE_DIR, f"{doc_id}.txt"))
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    doc = Document(id=doc_id, filename=filename, path=path, type="txt",
                   department=department, category=category, content=content)
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc
