"""Deliverables: PDF / DOCX / TXT / XLSX / PPTX. Includes actual generated answer."""
import html
import os
import re

from app.core.config import settings


def _out_path(task_id: str, ext: str) -> str:
    os.makedirs(settings.OUTPUT_DIR, exist_ok=True)
    name = f"Adrestia_Report_{task_id}.{ext}"
    return os.path.abspath(os.path.join(settings.OUTPUT_DIR, name))


def generate_pdf(task_id: str, title: str, content: str = "", findings: list[str] | None = None,
                 evidence: list[dict] | None = None, recommendation: str = "") -> str:
    """Generate a clean, professional PDF deliverable containing the actual generated answer."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

    path = _out_path(task_id, "pdf")
    doc = SimpleDocTemplate(
        path,
        pagesize=letter,
        rightMargin=48,
        leftMargin=48,
        topMargin=48,
        bottomMargin=48,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=10,
    )
    h2_style = ParagraphStyle(
        "DocH2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#1e293b"),
        spaceBefore=12,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "DocBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=15,
        textColor=colors.HexColor("#334155"),
        spaceAfter=6,
    )
    bullet_style = ParagraphStyle(
        "DocBullet",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#334155"),
        leftIndent=16,
        spaceAfter=4,
    )

    story = []
    story.append(Paragraph(html.escape(title), title_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=14))

    # Actual Generated Answer Content
    text_to_render = content.strip()
    if not text_to_render and findings:
        text_to_render = "\n".join([f"- {f}" for f in findings])

    if text_to_render:
        story.append(Paragraph("<b>Analysis & Generated Response</b>", h2_style))
        for line in text_to_render.splitlines():
            s = line.strip()
            if not s:
                story.append(Spacer(1, 4))
                continue
            # Replace arrows and special unicode for ReportLab standard fonts
            s_clean = s.replace("\u2192", "->").replace("→", "->")
            p_escaped = html.escape(s_clean)
            p_html = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', p_escaped)
            p_html = re.sub(r'\*([^*]+)\*', r'<i>\1</i>', p_html)

            if s_clean.startswith(("-", "*", "•")):
                cleaned_bullet = re.sub(r'^[-*•]\s*', '', p_html)
                story.append(Paragraph(f"&bull; {cleaned_bullet}", bullet_style))
            elif re.match(r'^\d+[\.\)]\s+', s_clean):
                story.append(Paragraph(p_html, bullet_style))
            elif s_clean.startswith("###"):
                heading_text = re.sub(r'^#{1,6}\s*', '', p_html)
                story.append(Paragraph(f"<b>{heading_text}</b>", h2_style))
            else:
                story.append(Paragraph(p_html, body_style))
        story.append(Spacer(1, 10))

    if evidence:
        story.append(Spacer(1, 6))
        story.append(Paragraph("<b>Grounding & Evidence Sources</b>", h2_style))
        for e in evidence:
            fname = html.escape(str(e.get("filename", "Document")))
            page = html.escape(str(e.get("page") or "Referenced"))
            excerpt = html.escape(str(e.get("excerpt", ""))[:140]).replace("\u2192", "->")
            story.append(Paragraph(f"&bull; <b>{fname}</b> ({page}): {excerpt}", bullet_style))

    doc.build(story)
    return path


def generate_docx(task_id: str, title: str, content: str = "", findings: list[str] | None = None,
                  evidence: list[dict] | None = None, recommendation: str = "") -> str:
    """Generate Word DOCX deliverable containing the actual generated response."""
    import docx

    path = _out_path(task_id, "docx")
    doc = docx.Document()
    doc.add_heading(title, level=1)

    text_to_render = content.strip()
    if not text_to_render and findings:
        text_to_render = "\n".join([f"- {f}" for f in findings])

    if text_to_render:
        doc.add_heading("Analysis & Generated Response", level=2)
        for line in text_to_render.splitlines():
            s = line.strip()
            if not s:
                continue
            is_bullet = s.startswith(("-", "*", "•"))
            is_numbered = bool(re.match(r'^\d+[\.\)]\s+', s))
            is_h3 = s.startswith("###")

            if is_h3:
                doc.add_heading(s.lstrip("# ").strip(), level=3)
                continue

            if is_bullet or is_numbered:
                p = doc.add_paragraph(style="List Bullet")
                raw_text = re.sub(r'^[-*•\d\.\)]\s*', '', s)
            else:
                p = doc.add_paragraph()
                raw_text = s

            parts = re.split(r'(\*\*[^*]+\*\*)', raw_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**") and len(part) >= 4:
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

    if evidence:
        doc.add_heading("Grounding & Evidence Sources", level=2)
        for e in evidence:
            fname = str(e.get("filename", "Document"))
            page = str(e.get("page") or "Referenced")
            excerpt = str(e.get("excerpt", ""))[:150]
            doc.add_paragraph(f"{fname} ({page}): {excerpt}", style="List Bullet")

    doc.save(path)
    return path


def generate_txt(task_id: str, title: str, content: str = "", findings: list[str] | None = None,
                 evidence: list[dict] | None = None, recommendation: str = "") -> str:
    path = _out_path(task_id, "txt")
    lines = [title, "=" * len(title), ""]
    text_to_render = content.strip()
    if not text_to_render and findings:
        text_to_render = "\n".join([f"- {f}" for f in findings])
    if text_to_render:
        lines += ["Analysis & Generated Response:", "-" * 32, text_to_render, ""]
    if evidence:
        lines += ["Evidence Sources:", "-" * 17]
        for e in evidence:
            lines.append(f"- {e.get('filename', '?')}, Page {e.get('page', '?')}: {e.get('excerpt', '')[:120]}")
        lines.append("")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return path


def generate_xlsx(task_id: str, findings: list[str], evidence: list[dict]) -> str:
    import openpyxl

    path = _out_path(task_id, "xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Findings"
    ws.append(["#", "Finding", "Evidence"])
    for i, f in enumerate(findings, 1):
        ev = evidence[i - 1] if i - 1 < len(evidence) else {}
        ws.append([i, f, f"{ev.get('filename', '')} p.{ev.get('page', '')}"])
    wb.save(path)
    return path


def generate_pptx(task_id: str, title: str, findings: list[str]) -> str:
    from pptx import Presentation
    from pptx.util import Pt

    path = _out_path(task_id, "pptx")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    for i, f in enumerate(findings, 1):
        box = slide.shapes.add_textbox(Pt(36), Pt(120 + i * 60), Pt(640), Pt(50))
        box.text_frame.text = f"{i}. {f}"
    prs.save(path)
    return path
